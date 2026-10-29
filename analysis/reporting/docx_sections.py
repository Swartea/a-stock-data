"""docx_sections — Section Registry 灰度块（P2-E S4 迁入，本轮新建）

本模块按 `docs/08-规范整改-P2E-文档渲染拆分-plan.md` §六 S4，从
`analysis/md_to_docx.py` **原样搬入**一段（§2.2 区块 7）：

| 区块 | 迁入符号 | 符号行 |
|---|---|---|
| 7 Section Registry 辅助 | `_load_result_near` `_render_md_block` | 70 |

`70` 与 §六 S4 行「搬 **70** 符号行」逐字对齐。口径见 §2.1：顶层符号的
`end_lineno - lineno + 1`，装饰器/docstring/注释/**块内空行**一并计入；
本文件自身的 import 块与文件头**不计入**迁入量（§五 明示）。
逐符号核算：`_load_result_near` 21 + `_render_md_block` 49 = **70** ✅

本片必须成立的不变量
--------------------
1. **函数体逐字未改** —— 只做「搬移 + 跨模块 import」，不做任何行为面清理。
   §九 第 7 条**不去重** `_render_md_block` 与主流程的 `while` 循环（它俩是同一套
   Markdown 行解析器的两份拷贝，docstring L180 自述「逻辑等价」）—— 去重会改
   解析分支顺序（R3/R4），是结构重设计，另立议题。
2. **扇入面已全部显式化** —— 区块 7 是 §三 里**耦合最深**的一块（扇出横跨 4 个区块）：
   对 `docx_primitives` 的 2 个依赖（`_fmt` `C_MUTED`）、对 `docx_paragraphs` 的
   3 个依赖（`add_heading` `add_quote` `add_body`）、对 `docx_tables` 的 1 个依赖
   （`flush_table`），以及对 `python-docx` 的 2 个依赖（`WD_ALIGN_PARAGRAPH` `Pt`）
   与标准库 2 个（`re` `Path`），全部走**裸名 import**，与薄壳的
   `from reporting.docx_primitives import ...` / `from sections import ...` 同一前提
   （`analysis/` 在 sys.path 上：脚本模式 `sys.path[0]`、`artifact_writer.py:108`
   显式 insert `_ANALYSIS_DIR`、CI `PYTHONPATH`）。
3. **零 `sections` 依赖** —— H4 的「裸名 + try 降级」只对 `sections` 这个**可选**
   依赖需要。区块 7 全文检索 `sections` **零命中**（实测）：`enabled_sections` 的
   消费点在**薄壳**区块 8 的灰度循环里，不在本模块。本模块不需要那层 try 降级；
   若将来有人顺手加进来又没套 try，CI 矩阵会 ImportError（R7）——
   本文件 `tests/reporting/test_docx_sections.py` 对此有可执行断言。
4. **薄壳 re-export** —— 这 **2 个名字**（2 个符号）在 `analysis/md_to_docx.py`
   顶层**原样再导出**，使 S4 之前的老调用路径继续可用（H2/H3）。两个在薄壳内
   **均仍有调用点**（`md_to_docx` 的 `_load_result_near(md_path)` 与
   `_render_md_block(doc, sec_md, stats)` 各 1 处，在灰度循环内），因此这条 import
   既是兼容面也是薄壳自身的依赖面；仍挂 `# noqa: F401` 以免 S5 收官时把这两个名字
   当成漏删而「顺手清理」。
5. **灰度块的隔离契约逐字未改** —— `_render_md_block` 末尾的
   `if in_table: flush_table(...)`（收尾表格）与 `_load_result_near` 的
   `except Exception: return None`（任何 IO/JSON 错误降级）都是**行为面**，一个字
   都不能动；`_load_result_near` 的两段式候选匹配（先按 stem 末尾 `-HHMM`
   时间戳精确匹配，再退回 `result_v3-*.json` 全量取 mtime 最新）同样是契约。
"""
import re
from pathlib import Path

from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt
from reporting.docx_paragraphs import add_body, add_heading, add_quote
from reporting.docx_primitives import C_MUTED, _fmt
from reporting.docx_tables import flush_table


# ---------------------------------------------------------------- Section Registry 辅助
def _load_result_near(md_path):
    """从 md_path 同目录读 result_v3-*.json (Task 1.5: 不依赖 V3 caller 改签名)
    Returns: dict 或 None (任何 IO/JSON 错误均返回 None)
    """
    try:
        import json as _json
        d = Path(md_path).parent
        # 优先匹配 stem 末尾 -HHMM 时间戳的 json
        stem = Path(md_path).stem  # e.g. "600693-东百集团-1616"
        cand = []
        m = re.search(r"-(\d{4})$", stem)
        if m:
            cand = list(d.glob(f"result_v3-{m.group(1)}.json"))
        if not cand:
            cand = list(d.glob("result_v3-*.json"))
        if not cand:
            return None
        latest = max(cand, key=lambda p: p.stat().st_mtime)
        return _json.loads(latest.read_text(encoding="utf-8"))
    except Exception:  # noqa: BLE001
        return None


def _render_md_block(doc, md_text, stats):
    """把一段 markdown 文本追加到 docx, 复用样式函数 (Task 1.5: section registry)
    逻辑与主流程的 while 循环等价, 但自带 cur_h2 / risk_zone / table_rows 状态
    """
    lines = md_text.split("\n")
    i = 0
    in_table = False
    table_rows = []
    cur_h2 = ""
    risk_zone = False
    while i < len(lines):
        line = lines[i]
        i += 1
        if line.strip().startswith("|") and line.strip().endswith("|"):
            table_rows.append(line)
            in_table = True
            continue
        elif in_table:
            flush_table(doc, table_rows, stats)
            in_table = False
        m = re.match(r"^(#{1,6})\s+(.+)$", line)
        if m:
            level = len(m.group(1))
            text = m.group(2).strip()
            add_heading(doc, text, level, stats)
            if level == 2:
                cur_h2 = text
                risk_zone = bool(re.search(r"风险警报", cur_h2))
            continue
        if re.match(r"^-{3,}$", line.strip()):
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            r = p.add_run("─" * 40)
            _fmt(r, size=9, color=C_MUTED)
            p.paragraph_format.space_before = Pt(4)
            p.paragraph_format.space_after = Pt(4)
            stats["paragraphs"] += 1
            continue
        if line.startswith(">"):
            add_quote(doc, line[1:], stats, in_conclusion=(cur_h2.startswith("🎯")))
            continue
        if not line.strip():
            continue
        if line.startswith("⚠️ 免责声明") or line.startswith("⚠ 免责声明"):
            add_body(doc, line, stats, disclaimer=True)
            continue
        add_body(doc, line, stats, risk_zone=risk_zone)
    if in_table:
        flush_table(doc, table_rows, stats)
