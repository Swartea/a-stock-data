"""docx_primitives — DOCX 渲染器的无业务语义地基（P2-E S1 迁入，本轮新建）

本模块按 `docs/08-规范整改-P2E-文档渲染拆分-plan.md` §六 S1，从
`analysis/md_to_docx.py` **原样搬入**三段（§2.2 区块 2/3/4）：

| 区块 | 迁入符号 | 符号行 |
|---|---|---|
| 2 设计 token | 13 行常量赋值（**15 个常量名**：`HX_RED`/`HX_GREEN`/`HX_AMBER` 同行三合一，加 `HX_ACCENT`）+ `_EMOJI_RE` | 14 |
| 3 XML 小工具 | `_set_fonts` `_fmt` `_ensure` `p_border` `cell_shade` `table_borders` `cell_bottom_border` `header_repeat` | 72 |
| 4 内容运行拆分 | `md_runs` `add_runs` | 22 |
| **合计** | | **108** |

`108` 与 §六 S1 行「搬 **108** 符号行」逐字对齐。口径见 §2.1：顶层符号的
`end_lineno - lineno + 1`，装饰器/docstring/注释/**块内空行**一并计入；
本文件自身的 import 块与文件头**不计入迁入量**（§五 明示）。

本片必须成立的不变量
--------------------
1. **函数体逐字未改** —— 只做「搬移 + 跨模块 import」，不做任何行为面清理。
   §十一 点名的两个死符号 `_ensure`(8 符号行) / `HX_ACCENT`(1 符号行) **原样保留**;
   删除属行为面清理，混进结构拆分会让「行为不变」这个判据失效，另开单片
   （§九 第 6 条 / §八 R6）。
2. **零业务语义** —— 本模块不含任何业务阈值或用户可见文案分支（§三：扇出 = 无/低）。
   3 处业务阈值（60% 数字列判定 / 「价位·数值」列 / 「综合」行加粗）全在区块 6，归 S2；
   全部中文关键词分支（`看空|看多`、`数据来源`、`⚠`、`🎯`、`风险警报`）在区块 5/7，归 S3/S4。
3. **零 `sections` 依赖** —— H4 的「裸名 + try 降级」只在薄壳与 `docx_sections` 侧需要。
   本模块只依赖 `re` 与 `python-docx`，是纯叶子层，S2-S4 可安全 import 它。
4. **薄壳 re-export** —— 这 **26 个名字**（24 个符号）在 `analysis/md_to_docx.py`
   顶层**原样再导出**，使 S1 之前的老调用路径（如 `from md_to_docx import H_HEADER`）
   继续可用（H2/H3）。名字数多于符号数，因区块 2 的 `HX_RED`/`HX_GREEN`/`HX_AMBER`
   三名同行赋值。
"""
import re

from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor

# ---------------------------------------------------------------- 设计 token
# docs/06 §2.1：禁纯黑；打印 AA 用加深过的警告/正色
C_DARK   = RGBColor(0x1C, 0x19, 0x17)   # 文字
C_ACCENT = RGBColor(0x25, 0x63, 0xEB)   # 强调（H2）
C_RED    = RGBColor(0xDC, 0x26, 0x26)   # 危险 / A股涨
C_GREEN  = RGBColor(0x15, 0x80, 0x3D)   # A股跌（#16a34a 加深到 AA）
C_AMBER  = RGBColor(0xB4, 0x53, 0x09)   # 警告（#d97706 加深到 AA）
C_MUTED  = RGBColor(0x57, 0x53, 0x4E)   # 次级文字/页脚
H_HAIR   = "E7E5E4"                      # hairline 分隔
H_HEADER = "D6D3D1"                      # 表头下加粗线
FILL_ZEBRA = "F5F5F4"
HX_RED, HX_GREEN, HX_AMBER = "DC2626", "15803D", "B45309"
HX_ACCENT = "2563EB"

F_CJK  = "Microsoft YaHei"   # 中文（微信/Word 通用）
F_MONO = "Consolas"          # 数字等宽

# emoji 拆段（只拆内容正文里出现的符号，标题里的不动；变体选择符 FE0F 并入同一 token，防止孤立 VS16）
_EMOJI_RE = re.compile("([\U0001F000-\U0001FAFF][️]?|[☀-➿][️]?|[⬀-⯿][️]?)")


# ---------------------------------------------------------------- XML 小工具
def _set_fonts(run, latin, cjk):
    """同时设拉丁与东亚字体"""
    run.font.name = latin
    run._element.rPr.rFonts.set(qn("w:eastAsia"), cjk)


def _fmt(run, size=None, bold=None, color=None, mono=False, italic=None):
    """统一字体格式入口：默认 微软雅黑/常规黑字"""
    _set_fonts(run, F_MONO if mono else F_CJK, F_CJK)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.font.bold = bold
    if italic is not None:
        run.font.italic = italic
    if color is not None:
        run.font.color.rgb = color
    return run


def _ensure(parent, tag, successors):
    """按 OOXML 顺序插子元素（succ=必须排在后面的标签名列表）"""
    el = OxmlElement(tag)
    try:
        parent.insert_element_before(el, *successors)
    except Exception:
        parent.append(el)
    return el


def p_border(p, edge="bottom", color=H_HEADER, sz=6, space=2):
    """段落边框（hairline）"""
    pPr = p._p.get_or_add_pPr()
    pBdr = pPr.find(qn("w:pBdr"))
    if pBdr is None:
        pBdr = OxmlElement("w:pBdr")
        pPr.insert_element_before(pBdr, "w:shd", "w:tabs", "w:spacing", "w:ind", "w:jc", "w:rPr")
    e = OxmlElement("w:" + edge)
    for k, v in (("w:val", "single"), ("w:sz", str(sz)), ("w:space", str(space)), ("w:color", color)):
        e.set(qn(k), v)
    pBdr.append(e)


def cell_shade(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = tcPr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tcPr.insert_element_before(shd, "w:noWrap", "w:tcMar", "w:vAlign", "w:textDirection", "w:tcFitText", "w:hideMark")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)


def table_borders(t, color=H_HAIR, sz=4):
    tblPr = t._tbl.tblPr
    borders = tblPr.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tblPr.insert_element_before(borders, "w:shd", "w:tblLayout", "w:tblCellMar", "w:tblLook", "w:tblCaption", "w:tblDescription")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        e = borders.find(qn("w:" + edge))
        if e is None:
            e = OxmlElement("w:" + edge)
            borders.append(e)
        for k, v in (("w:val", "single"), ("w:sz", str(sz)), ("w:space", "0"), ("w:color", color)):
            e.set(qn(k), v)


def cell_bottom_border(cell, color=H_HEADER, sz=6):
    tcPr = cell._tc.get_or_add_tcPr()
    tcB = tcPr.find(qn("w:tcBorders"))
    if tcB is None:
        tcB = OxmlElement("w:tcBorders")
        tcPr.insert_element_before(tcB, "w:shd", "w:noWrap", "w:tcMar", "w:vAlign", "w:textDirection", "w:tcFitText", "w:hideMark")
    e = OxmlElement("w:bottom")
    for k, v in (("w:val", "single"), ("w:sz", str(sz)), ("w:space", "0"), ("w:color", color)):
        e.set(qn(k), v)
    tcB.append(e)


def header_repeat(row):
    trPr = row._tr.get_or_add_trPr()
    th = OxmlElement("w:tblHeader")
    th.set(qn("w:val"), "true")
    trPr.append(th)


# ---------------------------------------------------------------- 内容运行拆分
def md_runs(content):
    """把 **加粗** 与 emoji 拆成 (text, bold, is_emoji) 元组流"""
    for seg in re.split(r"(\*\*[^*]+\*\*)", content):
        if not seg:
            continue
        bold = seg.startswith("**") and seg.endswith("**")
        if bold:
            seg = seg[2:-2]
        pos = 0
        for m in _EMOJI_RE.finditer(seg):
            if m.start() > pos:
                yield seg[pos:m.start()], bold, False
            yield m.group(1), bold, True
            pos = m.end()
        if pos < len(seg):
            yield seg[pos:], bold, False


def add_runs(p, content, size=10.5, color=None, bold_all=False, mono_all=False):
    """往段落追加内容运行；color 非空时非 emoji 文字上色，emoji 保持本色"""
    for text, bold, emoji in md_runs(content):
        r = p.add_run(text)
        rcol = None if (emoji or color is None) else color
        _fmt(r, size=size, bold=(bold_all or bold), color=rcol, mono=mono_all)
