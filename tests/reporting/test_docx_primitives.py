"""P2-E S1 · analysis/reporting/docx_primitives.py 的单元测试

只测**迁入的 24 个符号**（区块 2 设计 token + 区块 3 XML 小工具 + 区块 4 运行拆分），
**不测主流程**（§六 S1 第 2 步的边界）。`md_to_docx()` 本体与页脚 / Section Registry
灰度块分别属于区块 8 / S4，不在本文件断言范围内。

本模块存在的原因：S0 只锁住了「DOCX **产物**的版式语义」（语义 golden 投影），但
**迁入符号自身的取值与行为**没有任何直接断言 —— 拆分时若某个 token 抄错、某个
`insert_element_before` 的 successors 列表少一个标签，投影 golden 会红，但红在哪
一句说不清。这里把每个迁入符号的契约单独钉住，让 S1 的 diff 可被逐项复核。

零网络 / 零 live / 零外部数据：全部在内存里用 `python-docx` 建对象，不落盘、不读
`reports/`、不依赖当日跑批，也不读同目录 `result_v3-*.json`。

覆盖清单（24 个符号）：
- 区块 2（16 个名字 / 14 个符号）：C_DARK C_ACCENT C_RED C_GREEN C_AMBER C_MUTED
  H_HAIR H_HEADER FILL_ZEBRA HX_RED HX_GREEN HX_AMBER HX_ACCENT F_CJK F_MONO
  _EMOJI_RE
- 区块 3（8 个）：_set_fonts _fmt _ensure p_border cell_shade table_borders
  cell_bottom_border header_repeat
- 区块 4（2 个）：md_runs add_runs
另加一组**薄壳 re-export 契约**（H2/H3）：S1 之前的顶层可导入名字必须逐名仍在。

死符号保留（P2E §十一 / §八 R6 / §九 第 6 条）：`_ensure` 与 `HX_ACCENT` 全仓零调用
点，但本计划**一律不删**（删除属行为面清理，另开单片）。下面的
`test_dead_symbols_are_preserved` 把这条纪律变成可执行断言，防止后续片顺手删掉。
"""
import re
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[2] / "analysis"))

from docx import Document  # noqa: E402  # analysis/ 靠上面的 sys.path 引导
from docx.oxml import OxmlElement  # noqa: E402
from docx.oxml.ns import qn  # noqa: E402
from reporting import docx_primitives as prim  # noqa: E402


# ============================================================
# 工具
# ============================================================
def _rgb(run):
    """run 颜色; 未设色时 python-docx 抛 ValueError, 归一为 None。"""
    try:
        return run.font.color.rgb
    except (AttributeError, ValueError):
        return None


def _find(parent, tag):
    return parent.find(qn(tag))


def _find_all(parent, tag):
    return parent.findall(qn(tag))


def _attr(el, name):
    return el.get(qn(f"w:{name}"))


def _new_doc():
    return Document()


# ============================================================
# 区块 2 · 设计 token
# ============================================================
def test_color_tokens_are_the_frozen_rgb_values():
    """6 个 RGBColor token 取值逐个钉死（A 股红涨绿跌是用户可见约定，不可漂）。"""
    assert str(prim.C_DARK) == "1C1917"
    assert str(prim.C_ACCENT) == "2563EB"
    assert str(prim.C_RED) == "DC2626"
    assert str(prim.C_GREEN) == "15803D"
    assert str(prim.C_AMBER) == "B45309"
    assert str(prim.C_MUTED) == "57534E"


def test_hex_string_tokens_are_frozen():
    assert prim.H_HAIR == "E7E5E4"
    assert prim.H_HEADER == "D6D3D1"
    assert prim.FILL_ZEBRA == "F5F5F4"
    assert (prim.HX_RED, prim.HX_GREEN, prim.HX_AMBER) == ("DC2626", "15803D", "B45309")


def test_accent_hex_token_is_preserved_despite_zero_call_sites():
    """HX_ACCENT 是 §十一 点名的死符号: 零调用点, 但**一律不删**（§八 R6）。"""
    assert prim.HX_ACCENT == "2563EB"


def test_font_family_tokens():
    assert prim.F_CJK == "Microsoft YaHei"
    assert prim.F_MONO == "Consolas"


def test_emoji_re_splits_single_emoji_and_merges_variation_selector():
    """emoji 拆段: FE0F 变体选择符并入同一 token, 防止孤立 VS16（迁入前的原注释）。"""
    # U+1F3AF 🎯 落在 \U0001F000-\U0001FAFF 段
    assert [m.group(1) for m in prim._EMOJI_RE.finditer("a🎯b")] == ["🎯"]
    # U+26A0 ⚠ + U+FE0F 必须是一个 token, 不是两个
    assert [m.group(1) for m in prim._EMOJI_RE.finditer("⚠️x")] == ["⚠️"]
    assert [m.group(1) for m in prim._EMOJI_RE.finditer("无 emoji")] == []


# ============================================================
# 区块 3 · XML 小工具
# ============================================================
def test_set_fonts_sets_latin_and_east_asia():
    p = _new_doc().add_paragraph()
    r = p.add_run("x")
    prim._set_fonts(r, "Arial", "SimSun")
    assert r.font.name == "Arial"
    assert r._element.rPr.rFonts.get(qn("w:eastAsia")) == "SimSun"


def test_fmt_defaults_to_cjk_without_touching_size_bold_italic_color():
    p = _new_doc().add_paragraph()
    r = p.add_run("x")
    assert prim._fmt(r) is r, "_fmt 必须原样返回 run（调用方靠它链式用）"
    assert r.font.name == prim.F_CJK
    assert r._element.rPr.rFonts.get(qn("w:eastAsia")) == prim.F_CJK
    assert r.font.size is None
    assert r.font.bold is None
    assert r.font.italic is None
    assert _rgb(r) is None


def test_fmt_applies_size_bold_italic_color_and_mono():
    p = _new_doc().add_paragraph()
    r = p.add_run("x")
    prim._fmt(r, size=11, bold=True, italic=True, color=prim.C_RED, mono=True)
    assert r.font.name == prim.F_MONO, "mono=True 时拉丁字体走 F_MONO"
    assert r._element.rPr.rFonts.get(qn("w:eastAsia")) == prim.F_CJK, "东亚字体恒为 F_CJK"
    assert r.font.size.pt == 11
    assert r.font.bold is True
    assert r.font.italic is True
    assert str(_rgb(r)) == "DC2626"


def test_fmt_keeps_ea_asia_cjk_even_when_mono():
    """数字列「右对齐等宽」只换拉丁字体, 东亚字体不能跟着变（否则中文回落默认字体）。"""
    p = _new_doc().add_paragraph()
    r = p.add_run("中")
    prim._fmt(r, size=10.5, mono=True)
    assert r.font.name == prim.F_MONO
    assert r._element.rPr.rFonts.get(qn("w:eastAsia")) == prim.F_CJK


def test_ensure_inserts_before_first_matching_successor():
    p = _new_doc().add_paragraph()
    pPr = p._p.get_or_add_pPr()
    pPr.append(OxmlElement("w:jc"))
    el = prim._ensure(pPr, "w:pBdr", ("w:shd", "w:jc"))
    assert el.tag == qn("w:pBdr")
    kids = [k.tag for k in pPr]
    assert kids.index(qn("w:pBdr")) < kids.index(qn("w:jc"))


def test_ensure_falls_back_to_append_when_parent_cannot_insert():
    """降级分支: 父对象没有 insert_element_before 时走 append（§4 注释所述语义）。

    用 list 当父对象而不是真 lxml 元素: 真元素上这条 except 很难稳定触发,
    而这里要验的正是「AttributeError -> append」这条降级路径本身。
    """

    class _NoInsertElementBefore(list):
        pass

    parent = _NoInsertElementBefore()
    el = prim._ensure(parent, "w:pBdr", ("w:shd",))
    assert el.tag == qn("w:pBdr")
    assert parent == [el], "降级路径必须 append 到父对象"


def test_p_border_default_edge_and_values():
    """默认 edge=bottom, color=H_HEADER, sz=6, space=2。"""
    doc = _new_doc()
    p = doc.add_paragraph()
    prim.p_border(p)
    pBdr = _find(p._p.get_or_add_pPr(), "w:pBdr")
    e = _find(pBdr, "w:bottom")
    assert e is not None
    assert _attr(e, "val") == "single"
    assert _attr(e, "sz") == "6"
    assert _attr(e, "space") == "2"
    assert _attr(e, "color") == prim.H_HEADER


def test_p_border_custom_edge_and_reuses_single_pbdr_container():
    doc = _new_doc()
    p = doc.add_paragraph()
    prim.p_border(p, edge="bottom", color=prim.H_HAIR, sz=8, space=4)
    prim.p_border(p, edge="left", color=prim.HX_ACCENT, sz=18, space=4)
    pPr = p._p.get_or_add_pPr()
    assert len(_find_all(pPr, "w:pBdr")) == 1, "两次调用必须复用同一个 w:pBdr"
    pBdr = _find(pPr, "w:pBdr")
    left = _find(pBdr, "w:left")
    assert _attr(left, "sz") == "18"
    assert _attr(left, "color") == prim.HX_ACCENT


def test_cell_shade_sets_val_color_fill_and_is_idempotent():
    doc = _new_doc()
    t = doc.add_table(rows=1, cols=1)
    cell = t.rows[0].cells[0]
    prim.cell_shade(cell, prim.FILL_ZEBRA)
    tcPr = cell._tc.get_or_add_tcPr()
    assert len(_find_all(tcPr, "w:shd")) == 1
    shd = _find(tcPr, "w:shd")
    assert shd.get(qn("w:val")) == "clear"
    assert shd.get(qn("w:color")) == "auto"
    assert shd.get(qn("w:fill")) == prim.FILL_ZEBRA
    prim.cell_shade(cell, "FFFFFF")
    assert _find(cell._tc.get_or_add_tcPr(), "w:shd").get(qn("w:fill")) == "FFFFFF"


def test_table_borders_writes_exactly_six_edges_and_is_idempotent():
    doc = _new_doc()
    t = doc.add_table(rows=2, cols=2)
    prim.table_borders(t)
    prim.table_borders(t)  # 幂等: 不应变成 12 条边
    tblPr = t._tbl.tblPr
    assert len(_find_all(tblPr, "w:tblBorders")) == 1
    borders = _find(tblPr, "w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        e = _find(borders, f"w:{edge}")
        assert e is not None, f"缺边 {edge}"
        assert e.get(qn("w:val")) == "single"
        assert e.get(qn("w:sz")) == "4"
        assert e.get(qn("w:space")) == "0"
        assert e.get(qn("w:color")) == prim.H_HAIR
    assert len(list(borders)) == 6


def test_table_borders_custom_color_and_size():
    doc = _new_doc()
    t = doc.add_table(rows=1, cols=1)
    prim.table_borders(t, color=prim.H_HEADER, sz=12)
    borders = _find(t._tbl.tblPr, "w:tblBorders")
    top = _find(borders, "w:top")
    assert top.get(qn("w:color")) == prim.H_HEADER
    assert top.get(qn("w:sz")) == "12"


def test_cell_bottom_border_default_color_and_space_zero():
    doc = _new_doc()
    t = doc.add_table(rows=1, cols=1)
    cell = t.rows[0].cells[0]
    prim.cell_bottom_border(cell)
    tcB = _find(cell._tc.get_or_add_tcPr(), "w:tcBorders")
    e = _find(tcB, "w:bottom")
    assert e.get(qn("w:val")) == "single"
    assert e.get(qn("w:sz")) == "6"
    assert e.get(qn("w:space")) == "0", "单元格边框 space 恒为 0（与段落边框不同）"
    assert e.get(qn("w:color")) == prim.H_HEADER


def test_header_repeat_marks_row_as_repeating_header():
    doc = _new_doc()
    t = doc.add_table(rows=2, cols=2)
    prim.header_repeat(t.rows[0])
    trPr = t.rows[0]._tr.get_or_add_trPr()
    assert len(_find_all(trPr, "w:tblHeader")) == 1
    assert _find(trPr, "w:tblHeader").get(qn("w:val")) == "true"


# ============================================================
# 区块 4 · 内容运行拆分
# ============================================================
def test_md_runs_plain_text_single_tuple():
    assert list(prim.md_runs("hello")) == [("hello", False, False)]


def test_md_runs_splits_bold_segments():
    assert list(prim.md_runs("a**b**c")) == [
        ("a", False, False),
        ("b", True, False),
        ("c", False, False),
    ]


def test_md_runs_marks_emoji_and_keeps_bold_flag():
    assert list(prim.md_runs("**hi 🎯**")) == [
        ("hi ", True, False),
        ("🎯", True, True),
    ]


def test_md_runs_empty_string_yields_nothing():
    assert list(prim.md_runs("")) == []


def test_add_runs_colors_text_but_not_emoji():
    """color 非空时**非 emoji** 文字上色, emoji 保持本色（迁入前的原注释）。"""
    p = _new_doc().add_paragraph()
    prim.add_runs(p, "价格 🎯", size=10, color=prim.C_RED)
    runs = p.runs
    assert [r.text for r in runs] == ["价格 ", "🎯"]
    assert str(_rgb(runs[0])) == "DC2626"
    assert _rgb(runs[1]) is None, "emoji 必须保持本色"


def test_add_runs_defaults_to_cjk_10_5pt():
    p = _new_doc().add_paragraph()
    prim.add_runs(p, "普通行")
    r = p.runs[0]
    assert r.font.size.pt == 10.5
    assert r.font.name == prim.F_CJK
    assert _rgb(r) is None
    # add_runs 恒传 bold=(bold_all or bold), 故 bold 是显式 bool 而非 None
    assert r.font.bold is False


def test_add_runs_bold_all_and_mono_all():
    p = _new_doc().add_paragraph()
    prim.add_runs(p, "无标记", bold_all=True, mono_all=True, size=11)
    r = p.runs[0]
    assert r.font.bold is True, "bold_all=True 时无 ** 标记也加粗"
    assert r.font.name == prim.F_MONO
    assert r.font.size.pt == 11


# ============================================================
# 薄壳 re-export 契约（H2 / H3）+ §六「禁止新旧两套实现并存」
# ============================================================
_MOVED_NAMES = {
    # 区块 2 设计 token（16 个名字 = 13 行赋值含 15 个常量名 + _EMOJI_RE）
    "_EMOJI_RE", "C_ACCENT", "C_AMBER", "C_DARK", "C_GREEN", "C_MUTED", "C_RED",
    "F_CJK", "F_MONO", "FILL_ZEBRA", "H_HAIR", "H_HEADER", "HX_ACCENT",
    "HX_AMBER", "HX_GREEN", "HX_RED",
    # 区块 3 XML 小工具（8 个）
    "_ensure", "_fmt", "_set_fonts", "cell_bottom_border", "cell_shade",
    "header_repeat", "p_border", "table_borders",
    # 区块 4 内容运行拆分（2 个）
    "add_runs", "md_runs",
}
_MOVED_FUNCS = (
    "_set_fonts", "_fmt", "_ensure", "p_border", "cell_shade",
    "table_borders", "cell_bottom_border", "header_repeat", "md_runs", "add_runs",
)
_MOVED_CONSTS = tuple(sorted(n for n in _MOVED_NAMES if n not in _MOVED_FUNCS and n != "_EMOJI_RE"))


def test_moved_surface_is_exactly_26_names():
    """迁出面自检: 24 个符号对应 26 个名字（HX_RED/HX_GREEN/HX_AMBER 同行三合一）。"""
    assert len(_MOVED_NAMES) == 26, "26 个名字（§六 S1: 15 常量名 + 9 叶子符号 + _EMOJI_RE + add_runs）"
    assert len(_MOVED_FUNCS) == 10, "区块 3 的 8 个 + 区块 4 的 2 个"
    assert len(_MOVED_CONSTS) == 15, "15 个常量名 / 13 行赋值（§六 S1 行）"


def test_shell_reexports_every_moved_name():
    """S1 之前的 26 个顶层名字必须逐名仍在薄壳上（H2 双人格 + H3 假模块面）。

    用 hasattr 而不是 `from md_to_docx import (...)`: 后者一旦缺名会在收集期直接
    ImportError, 且报错不告诉你是**哪个**名字丢了; hasattr 一次报全部缺失。
    """
    import md_to_docx as shell

    missing = sorted(n for n in _MOVED_NAMES if not hasattr(shell, n))
    assert not missing, f"薄壳丢失 re-export 名字: {missing}"


def test_shell_reexport_values_are_identical_objects_not_copies():
    """re-export 必须是同一对象而非复制品 —— 否则薄壳与新模块会各走各的样式值。"""
    import md_to_docx as shell

    assert shell.C_DARK is prim.C_DARK
    assert shell._fmt is prim._fmt
    assert shell._EMOJI_RE is prim._EMOJI_RE
    assert shell.table_borders is prim.table_borders


def test_shell_keeps_docx_import_names_for_backward_compat():
    """OxmlElement / RGBColor 迁出后薄壳无调用点, 但**为兼容保留**（2026-10-06 拍板）。"""
    import md_to_docx as shell

    assert hasattr(shell, "OxmlElement")
    assert hasattr(shell, "RGBColor")


def test_shell_no_longer_defines_moved_symbols_in_place(analysis_dir):
    """§六 硬规则「禁止新旧两套实现并存」: 薄壳不得再原地定义迁出符号。

    只查**定义**, 不查名字出现 —— re-export import 与调用点当然会出现这些名字。
    """
    src = (analysis_dir / "md_to_docx.py").read_text(encoding="utf-8")
    still_def = [n for n in _MOVED_FUNCS if re.search(rf"^def {n}\(", src, re.M)]
    assert not still_def, f"薄壳仍原地定义 {still_def} —— 新旧两套实现并存"
    still_assign = [n for n in _MOVED_CONSTS if re.search(rf"^{n}\s*=", src, re.M)]
    assert not still_assign, f"薄壳仍原地赋值 {still_assign} —— 新旧两套实现并存"
    assert not re.search(r"^_EMOJI_RE\s*=", src, re.M), "薄壳仍原地定义 _EMOJI_RE"
